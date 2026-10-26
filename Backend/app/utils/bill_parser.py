"""
Supplier-bill parser  (the "Extract" and "Transform" steps of the stock-import ETL)

Everything in here is plain Python: no database, no web framework.
It turns the different shapes a supplier bill can arrive in

    * a CSV / Excel table
    * the text of a digital PDF
    * OCR lines (with positions) from a photographed bill

into ONE common row format:

    {
        "line_no": 1,
        "item_name": "CLOPIGOLD A 150(CLOPI 75 & ASPRN 15)",
        "composition_hint": "CLOPI 75 & ASPRN 15",   # text in brackets, if any
        "code_hint": None,                           # PMBI code printed in brackets, if any
        "pmbi_code": None,                           # only when the file has a code column
        "hsn": "30049099",
        "batch_number": "CJSIK503",
        "expiry_date": "2027-10-31",                 # always the LAST day of the month
        "quantity": 1,                               # billed qty + free qty
        "free_qty": 0,
        "pack_size": "15's",
        "mrp": 48.0,
        "purchase_rate": None,
        "manufacturer": None,
        "amount": 48.0,
        "raw": "<the text the row was read from>",
        "warnings": [],
    }

Nothing here is trusted blindly: the review screen lets the owner fix every field.
"""

from __future__ import annotations

import calendar
import csv
import io
import math
import re
import statistics
from datetime import date, datetime
from typing import Any, Iterable, Optional


# =====================================================
# Small helpers
# =====================================================

MONTHS: dict[str, int] = {}
for _i in range(1, 13):
    MONTHS[calendar.month_abbr[_i].lower()] = _i
    MONTHS[calendar.month_name[_i].lower()] = _i
MONTHS["sept"] = 9


def _clean(value: Any) -> str:
    """Any cell value -> tidy single-line string."""
    if value is None:
        return ""
    if isinstance(value, float) and math.isnan(value):
        return ""
    return re.sub(r"\s+", " ", str(value).replace(" ", " ")).strip()


def parse_number(value: Any) -> Optional[float]:
    """'1,234.50', 'Rs. 48.00', 48 -> float.  Returns None if it is not a number."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return None if (isinstance(value, float) and math.isnan(value)) else float(value)
    text = _clean(value)
    text = re.sub(r"(?i)^(rs\.?|inr|₹)\s*", "", text)
    text = text.replace(",", "")
    if not re.fullmatch(r"-?\d+(\.\d+)?", text):
        return None
    return float(text)


def _to_int_qty(number: Optional[float]) -> tuple[Optional[int], Optional[str]]:
    """Quantity as a whole number (+ a warning when it was fractional)."""
    if number is None:
        return None, None
    if number < 0:
        return None, f"Negative quantity {number:g} ignored."
    rounded = int(round(number))
    if abs(number - rounded) > 1e-6:
        return rounded, f"Quantity {number:g} is not a whole number - rounded to {rounded}."
    return rounded, None


def _last_day(year: int, month: int) -> date:
    return date(year, month, calendar.monthrange(year, month)[1])


def _year(text: str) -> int:
    year = int(text)
    return 2000 + year if year < 100 else year


# =====================================================
# Expiry dates
# =====================================================

def parse_expiry(value: Any) -> Optional[date]:
    """
    Understands every way supplier bills print an expiry:

        Oct-2027  OCT 2027  Oct'27  10/27  10/2027  10-27  2027-10
        31/10/2027  2027-10-31  and real Excel dates

    A month + year means "good until the end of that month", so the result
    is the LAST day of that month.
    """
    if value is None:
        return None
    if isinstance(value, datetime):
        value = value.date()
    if isinstance(value, date):
        # Excel stores "Oct-2027" as 1 Oct 2027 - treat a 1st as month/year.
        return _last_day(value.year, value.month) if value.day == 1 else value

    text = _clean(value).upper().replace("’", "'").replace("`", "'")
    if not text:
        return None

    try:
        match = re.fullmatch(r"(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})", text)
        if match:
            return date(int(match[1]), int(match[2]), int(match[3]))

        match = re.fullmatch(r"(\d{1,2})[-/.](\d{1,2})[-/.](\d{2}|\d{4})", text)
        if match:
            return date(_year(match[3]), int(match[2]), int(match[1]))

        match = re.fullmatch(r"(\d{4})[-/.](\d{1,2})", text)
        if match and 1 <= int(match[2]) <= 12:
            return _last_day(int(match[1]), int(match[2]))

        match = re.fullmatch(r"(\d{1,2})\s*[-/.]\s*(\d{4}|\d{2})", text)
        if match and 1 <= int(match[1]) <= 12:
            return _last_day(_year(match[2]), int(match[1]))

        match = re.fullmatch(r"([A-Z]{3,9})\s*[-/.' ]\s*'?(\d{4}|\d{2})", text)
        if match and match[1].lower() in MONTHS:
            return _last_day(_year(match[2]), MONTHS[match[1].lower()])
    except ValueError:
        return None
    return None


# An expiry-looking token inside a line of text (validated by parse_expiry afterwards)
EXPIRY_IN_TEXT = re.compile(
    r"(?<![\w/.'-])("
    r"(?:0?[1-9]|1[0-2])\s?[/-]\s?(?:\d{4}|\d{2})"
    r"|[A-Za-z]{3,9}\s?[-/.' ]\s?'?(?:\d{4}|\d{2})"
    r"|\d{4}-(?:0[1-9]|1[0-2])(?:-\d{2})?"
    r")(?![\w/-])"
)


def find_expiry_in_text(text: str) -> Optional[tuple[int, int, date]]:
    """First valid expiry token that has some text in front of it -> (start, end, date)."""
    for match in EXPIRY_IN_TEXT.finditer(text):
        parsed = parse_expiry(match.group(1))
        if parsed and text[: match.start()].strip():
            return match.start(1), match.end(1), parsed
    return None


# =====================================================
# Pack size
# =====================================================

PACK_TOKEN = re.compile(r"(?i)^(\d{1,4})\s*['’`]?\s*s$|^(\d{1,3})\s*[xX]\s*(\d{1,3})(?:\s*[xX]\s*(\d{1,3}))?$")


def normalize_pack(text: Any) -> Optional[str]:
    """'15s', "15'S", '15 s' -> "15's";  '10x10' stays '10x10'."""
    value = _clean(text)
    if not value:
        return None
    match = PACK_TOKEN.match(value)
    if not match:
        return value
    if match.group(1):
        return f"{int(match.group(1))}'s"
    parts = [p for p in match.groups()[1:] if p]
    return "x".join(str(int(p)) for p in parts)


def looks_like_pack(token: str) -> bool:
    return bool(PACK_TOKEN.match(token.strip()))


# =====================================================
# Column headers  (CSV / Excel / PDF tables)
# =====================================================

# normalised header text -> column kind
HEADER_KINDS: dict[str, str] = {}


def _register(kind: str, names: Iterable[str]) -> None:
    for name in names:
        HEADER_KINDS[re.sub(r"[^a-z0-9]", "", name.lower())] = kind


_register("item", [
    "item", "item name", "itemname", "product", "product name", "description", "item description",
    "particulars", "medicine", "medicine name", "drug", "drug name", "name", "goods", "goods supplied",
    "name of item", "name of product", "article",
])
_register("code", ["pmbi code", "pmbicode", "pmbi", "drug code", "drugcode", "product code", "item code", "code"])
_register("hsn", ["hsn", "hsn sac", "hsn code", "sac"])
_register("batch", ["batch", "batch no", "batch number", "batch no.", "lot", "lot no", "batchno"])
_register("expiry", ["exp", "expiry", "exp dt", "exp date", "expiry date", "exp.", "expdt", "date of expiry"])
_register("qty", ["qty", "quantity", "qnty", "billed qty", "ship qty", "qty."])
_register("free", ["free", "free qty", "fr", "fr qty", "bonus", "scheme"])
_register("pack", ["pack", "packing", "pkg", "unit", "pack size", "uom", "packaging"])
_register("mrp", ["mrp", "m.r.p", "m r p", "list price", "max retail price", "mrp inr"])
_register("price", ["price"])
_register("rate", ["rate", "ptr", "pts", "purchase rate", "net rate", "p rate", "landing rate", "cost", "unit rate", "net price"])
_register("mfr", ["mfr", "mfg", "manufacturer", "company", "mkt", "mktd", "marketed by", "make", "manufactured by"])
_register("amount", ["amount", "amt", "net amount", "total", "value", "line amount"])
_register("tax", ["cgst", "sgst", "igst", "gst", "tax", "cgst amt", "sgst amt", "igst amt", "gst %", "disc", "discount", "disc %", "cd"])


def header_kind(header: Any) -> Optional[str]:
    """Which kind of column does this header text describe?  None if unknown."""
    key = re.sub(r"[^a-z0-9]", "", _clean(header).lower())
    if not key:
        return None
    if key in HEADER_KINDS:
        return HEADER_KINDS[key]
    if key.startswith("batch") or key.startswith("lot"):
        return "batch"
    if key.startswith("exp"):
        return "expiry"
    if key.startswith("qty") or key.startswith("quant"):
        return "qty"
    if key.startswith("free"):
        return "free"
    if "mrp" in key:
        return "mrp"
    if key.startswith("hsn"):
        return "hsn"
    if any(t in key for t in ("cgst", "sgst", "igst")) or key.startswith("disc"):
        return "tax"
    if key.startswith("amount") or key.endswith("amt"):
        return "amount"
    return None


# Lines that end the item table (totals, tax summary, signatures ...)
FOOTER_WORDS = re.compile(
    r"(?i)grand\s*total|sub\s*total|total\s*tax|tax\s*rate|taxable\s*am|rupees|round(?:ed)?\s*off|"
    r"terms|e\.?\s*&\s*o\.?\s*e|subject\s*to|declaration|signature|authori[sz]ed|no\s*exchange|"
    r"net\s*amount|bank\s*details|amount\s*in\s*words|certified|for\s+\w+.*\bltd\b"
)


# =====================================================
# Row assembly
# =====================================================

CATEGORY_CODE = re.compile(r"^[A-Z]{1,3}-[A-Z]{2,5}$")        # "G-CVS", "J-CVS", "G-ADC"
HSN_TOKEN = re.compile(r"^\d{4}(\d{2}(\d{2})?)?$")             # 3004 / 300490 / 30049099
PHARMA_HSN = re.compile(r"^30\d{2}(\d{2}(\d{2})?)?$")
PAREN_GROUP = re.compile(r"\(([^()]*)\)")


def _split_hints(item_name: str) -> tuple[Optional[str], Optional[str]]:
    """Text in brackets -> (composition_hint, pmbi code hint)."""
    composition, code = None, None
    for group in PAREN_GROUP.findall(item_name):
        group = group.strip()
        if re.fullmatch(r"\d{2,5}", group):
            code = group
        elif re.search(r"[A-Za-z]{3,}", group):
            composition = group
    return composition, code


def _new_row(line_no: int, raw: str) -> dict:
    return {
        "line_no": line_no, "item_name": "", "composition_hint": None, "code_hint": None,
        "pmbi_code": None, "hsn": None, "batch_number": None, "expiry_date": None,
        "quantity": None, "free_qty": 0, "pack_size": None, "mrp": None, "purchase_rate": None,
        "manufacturer": None, "amount": None, "raw": raw, "warnings": [],
    }


def _finish_row(row: dict) -> dict:
    """Common post-processing for every parsed row."""
    row["item_name"] = _clean(row["item_name"])
    composition, code = _split_hints(row["item_name"])
    row["composition_hint"] = composition
    row["code_hint"] = code
    if row["batch_number"]:
        row["batch_number"] = _clean(row["batch_number"]).strip(".,;:").upper()
    if not row["batch_number"]:
        row["batch_number"] = None
        row["warnings"].append("Batch number not found.")
    if not row["expiry_date"]:
        row["warnings"].append("Expiry date not found.")
    if not row["quantity"]:
        row["warnings"].append("Quantity not found.")
    # Cross-check: qty x price should equal the printed amount.
    prices = [p for p in (row["purchase_rate"], row["mrp"]) if p]
    if prices and row["amount"] and row["quantity"]:
        billed = max(row["quantity"] - (row["free_qty"] or 0), 1)
        tolerance = max(0.02 * row["amount"], 1.0)
        if not any(abs(qty * price - row["amount"]) <= tolerance
                   for price in prices for qty in (billed, row["quantity"])):
            row["warnings"].append("Quantity x price does not equal the printed amount - please check the numbers.")
    return row


# ----- tables (CSV / Excel / PDF tables) ---------------------------------------------

def _is_footer_cells(cells: list[str]) -> bool:
    first = next((c for c in cells if c), "")
    return bool(re.match(r"(?i)^(grand\s*)?total\b|^sub\s*total|^round|^net\s*(amount|payable)", first))


def rows_from_table(table: list[list[Any]]) -> tuple[list[dict], list[str], bool]:
    """
    A table whose first useful row is a header.
    Returns (rows, unparsed_lines, header_found).
    """
    kept = [row for row in table if any(_clean(c) for c in row)]
    grid = [[_clean(c) for c in row] for row in kept]
    raw_grid = [list(row) for row in kept]

    header_index, columns = None, {}
    for index, row in enumerate(grid[:20]):
        kinds = {}
        for col, cell in enumerate(row):
            kind = header_kind(cell)
            if kind and kind not in kinds.values():
                kinds[col] = kind
        has_core = ("batch" in kinds.values() or "expiry" in kinds.values()) and (
            "item" in kinds.values() or "code" in kinds.values()
        )
        if has_core and len(kinds) >= 3:
            header_index, columns = index, kinds
            break

    if header_index is None:
        return [], [], False

    # a bare "Price" column means MRP unless there is a real MRP column
    kinds_present = set(columns.values())
    for col, kind in list(columns.items()):
        if kind == "price":
            columns[col] = "rate" if "mrp" in kinds_present else "mrp"

    rows: list[dict] = []
    unparsed: list[str] = []
    for line_no, (cells, raw_cells) in enumerate(zip(grid[header_index + 1:], raw_grid[header_index + 1:]), start=1):
        joined = " ".join(c for c in cells if c)
        if _is_footer_cells(cells) or (FOOTER_WORDS.search(joined) and not any(
            cells[col] for col, kind in columns.items() if kind in ("batch", "expiry") and col < len(cells)
        )):
            continue
        row = _new_row(line_no, joined)
        qty_value, free_value = None, 0
        for col, kind in columns.items():
            if col >= len(cells):
                continue
            cell = cells[col]
            if not cell:
                continue
            if kind == "item":
                row["item_name"] = cell
            elif kind == "code":
                row["pmbi_code"] = cell.strip()
            elif kind == "hsn":
                row["hsn"] = cell
            elif kind == "batch":
                row["batch_number"] = cell
            elif kind == "expiry":
                raw_cell = raw_cells[col] if col < len(raw_cells) else cell
                parsed = parse_expiry(raw_cell if isinstance(raw_cell, (date, datetime)) else cell)
                row["expiry_date"] = parsed.isoformat() if parsed else None
                if not parsed:
                    row["warnings"].append(f"Could not read expiry '{cell}'.")
            elif kind == "qty":
                qty_value = parse_number(cell)
            elif kind == "free":
                free_value = parse_number(cell) or 0
            elif kind == "pack":
                row["pack_size"] = normalize_pack(cell)
            elif kind == "mrp":
                row["mrp"] = parse_number(cell)
            elif kind == "rate":
                row["purchase_rate"] = parse_number(cell)
            elif kind == "mfr":
                row["manufacturer"] = cell
            elif kind == "amount":
                row["amount"] = parse_number(cell)
        if not row["item_name"] and not row["pmbi_code"]:
            if joined:
                unparsed.append(joined)
            continue
        qty, warning = _to_int_qty(qty_value)
        free, _ = _to_int_qty(free_value)
        if warning:
            row["warnings"].append(warning)
        row["free_qty"] = free or 0
        row["quantity"] = (qty or 0) + row["free_qty"] if (qty or free) else None
        if row["pmbi_code"] and not row["item_name"]:
            row["item_name"] = f"PMBI {row['pmbi_code']}"
        rows.append(_finish_row(row))
    return rows, unparsed, True


def rows_from_csv_bytes(data: bytes) -> tuple[list[dict], list[str], bool]:
    for encoding in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeError:
            continue
    else:  # pragma: no cover
        text = data.decode("latin-1", errors="replace")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    table = list(csv.reader(io.StringIO(text), dialect))
    return rows_from_table(table)


def rows_from_xlsx_bytes(data: bytes) -> tuple[list[dict], list[str], bool]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("Excel import needs the 'openpyxl' package (pip install openpyxl).") from exc
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    for sheet in workbook.worksheets:
        table = [list(row) for row in sheet.iter_rows(values_only=True)]
        rows, unparsed, found = rows_from_table(table)
        if found:
            return rows, unparsed, True
    return [], [], False


# ----- text rows (digital PDF text / OCR) --------------------------------------------

# keywords that describe the column ORDER of a printed table header
HEADER_WORDS = [
    ("item", r"goods|particulars|description|item|product|medicine|drug"),
    ("hsn", r"hsn(?:\s*/\s*sac)?"),
    ("batch", r"batch(?:\s*no\.?)?|lot(?:\s*no\.?)?"),
    ("expiry", r"exp(?:iry)?(?:\s*d(?:t|ate))?\.?"),
    ("qty", r"qty\.?|quantity"),
    ("free", r"free(?:\s*qty)?"),
    ("pack", r"unit|pack(?:ing)?|pkg"),
    ("mrp", r"mrp|m\.r\.p|list\s*price|price"),
    ("rate", r"rate|ptr|pts"),
    ("tax", r"cgst|sgst|igst|gst|tax|disc(?:ount)?"),
    ("amount", r"amount|amt|total|value"),
]
_HEADER_RE = re.compile(
    "|".join(f"(?P<{kind}>\\b(?:{pattern})\\b)" for kind, pattern in HEADER_WORDS), re.IGNORECASE
)


def header_layout(text: str) -> list[str]:
    """
    'HSN/SAC Batch Exp Dt Qty. Unit List Price CGST Amt. SGST Amt. Amount'
        -> ['hsn','batch','expiry','qty','pack','mrp','tax','tax','amount']
    Only the ORDER matters (photos are tilted, so x-positions cannot be trusted).
    """
    layout: list[str] = []
    previous = None
    for match in _HEADER_RE.finditer(text):
        kind = match.lastgroup
        directly_after_tax = previous == "tax"
        previous = kind
        if kind == "amount" and directly_after_tax:
            continue                       # "CGST Amt." is one tax column, not an amount column
        if layout and layout[-1] == kind and kind != "tax":
            continue
        layout.append(kind)
    return layout


def looks_like_header(text: str) -> bool:
    kinds = set(header_layout(text))
    return len(kinds) >= 3 and ("batch" in kinds or "expiry" in kinds) and bool({"qty", "mrp", "rate", "amount"} & kinds)


def parse_text_row(text: str, layout: Optional[list[str]], line_no: int) -> Optional[dict]:
    """One printed bill line -> common row.  None when the line is not an item line."""
    found = find_expiry_in_text(text)
    if not found:
        return None
    start, end, expiry = found
    left, right = text[:start].strip(), text[end:].strip()

    layout = layout or []
    after_expiry_first = (
        layout.index("expiry") < layout.index("batch") if "expiry" in layout and "batch" in layout else False
    )

    row = _new_row(line_no, text)
    row["expiry_date"] = expiry.isoformat()

    left_tokens = left.split()
    right_tokens = right.split()

    # ---- batch: normally the token just before the expiry
    batch = None
    if after_expiry_first and right_tokens:
        batch = right_tokens.pop(0)
    elif left_tokens:
        batch = left_tokens.pop()
    if batch and not re.search(r"[A-Za-z0-9]", batch):
        batch = None
    row["batch_number"] = batch

    # ---- HSN code just before the batch
    if left_tokens and HSN_TOKEN.match(left_tokens[-1]) and (PHARMA_HSN.match(left_tokens[-1]) or "hsn" in layout):
        row["hsn"] = left_tokens.pop()

    # ---- leading serial number ("1", "12.")
    if len(left_tokens) > 1 and re.fullmatch(r"\d{1,3}\.?", left_tokens[0]):
        left_tokens.pop(0)
    row["item_name"] = " ".join(left_tokens)

    # ---- numbers after the expiry, read in the order the header printed them
    order = []
    if layout:
        anchor = layout.index("expiry") if "expiry" in layout else -1
        batch_pos = layout.index("batch") if "batch" in layout else -1
        order = [k for k in layout[max(anchor, batch_pos) + 1:] if k not in ("item", "hsn", "batch", "expiry")]
    if not order:
        order = ["qty", "pack", "mrp"]

    values: dict[str, Any] = {}
    cursor = 0
    for kind in order:
        if cursor >= len(right_tokens):
            break
        token = right_tokens[cursor]
        if kind == "pack":
            if looks_like_pack(token):
                values["pack"] = normalize_pack(token)
                cursor += 1
            continue
        number = parse_number(token)
        if number is None:
            if looks_like_pack(token) and "pack" not in values:
                values["pack"] = normalize_pack(token)       # pack printed where we did not expect it
                cursor += 1
                number = parse_number(right_tokens[cursor]) if cursor < len(right_tokens) else None
                if number is None:
                    continue
                token = right_tokens[cursor]
            else:
                continue
        cursor += 1
        if kind in ("qty", "free", "mrp", "rate", "amount") and kind not in values:
            values[kind] = number
        elif kind == "price":
            values.setdefault("mrp", number)
        # "tax" columns are read and thrown away

    qty, warning = _to_int_qty(values.get("qty"))
    free, _ = _to_int_qty(values.get("free"))
    if warning:
        row["warnings"].append(warning)
    row["free_qty"] = free or 0
    row["quantity"] = (qty or 0) + row["free_qty"] if (qty or free) else None
    row["pack_size"] = values.get("pack")
    row["mrp"] = values.get("mrp")
    row["purchase_rate"] = values.get("rate")
    row["amount"] = values.get("amount")
    return row


def _drop_category_codes(rows: list[dict]) -> None:
    """
    'G-CVS CLOPIGOLD ...' : drop the leading category code, but only when
    (nearly) every row has one - otherwise it is probably a real brand like D-RISE.
    """
    if len(rows) < 2:
        return
    first = [r["item_name"].split(" ", 1)[0] if r["item_name"] else "" for r in rows]
    coded = [bool(CATEGORY_CODE.match(t)) for t in first]
    if sum(coded) / len(rows) >= 0.6:
        for row, is_coded in zip(rows, coded):
            if is_coded and " " in row["item_name"]:
                row["item_name"] = row["item_name"].split(" ", 1)[1]


def rows_from_text_rows(text_rows: list[str]) -> tuple[list[dict], list[str], bool]:
    """
    `text_rows` = one string per printed line, top to bottom
    (digital-PDF text lines, or OCR boxes already grouped into lines).
    """
    layout: Optional[list[str]] = None
    header_found = False
    in_table = False
    rows: list[dict] = []
    unparsed: list[str] = []

    cleaned = [_clean(t) for t in text_rows]
    header_at = next((i for i, t in enumerate(cleaned) if looks_like_header(t)), None)
    if header_at is not None:
        layout = header_layout(cleaned[header_at])
        header_found = in_table = True
        cleaned = cleaned[header_at + 1:]          # invoice number, address, ... are not items

    for text in cleaned:
        if not text:
            continue
        if header_found and FOOTER_WORDS.search(text) and not find_expiry_in_text(text):
            in_table = False
            continue
        if header_found and not in_table:
            continue
        row = parse_text_row(text, layout, len(rows) + 1)
        if row:
            rows.append(row)
        elif header_found and re.search(r"[A-Za-z]{3,}", text):
            unparsed.append(text)

    _drop_category_codes(rows)
    for row in rows:
        _finish_row(row)
    return rows, unparsed, header_found


# ----- OCR lines with positions ------------------------------------------------------

def group_ocr_lines(lines: list[dict]) -> list[str]:
    """
    PaddleOCR gives separate boxes (text + position).  Put the boxes that sit on the
    same printed line back together, left to right.  Photos are usually tilted, so the
    tilt angle is estimated from the long text boxes and removed first.
    """
    items = []
    for line in lines:
        text = _clean(line.get("text"))
        box = line.get("box")
        if not text or not box or len(box) < 4:
            continue
        left, top, right, bottom = [float(v) for v in box[:4]]
        items.append({
            "text": text, "cx": (left + right) / 2, "cy": (top + bottom) / 2,
            "w": right - left, "h": max(bottom - top, 1.0), "angle": line.get("angle"),
        })
    if not items:
        return []

    angles = [i["angle"] for i in items if i["angle"] is not None and i["w"] > 4 * i["h"]]
    skew = statistics.median(angles) if len(angles) >= 3 else 0.0
    if abs(skew) > 15:          # nonsense - probably a rotated page; leave it alone
        skew = 0.0
    sin, cos = math.sin(math.radians(skew)), math.cos(math.radians(skew))
    for item in items:
        item["x"] = item["cx"] * cos + item["cy"] * sin
        item["y"] = -item["cx"] * sin + item["cy"] * cos

    median_height = statistics.median(i["h"] for i in items)
    tolerance = 0.6 * median_height

    items.sort(key=lambda i: i["y"])
    grouped: list[list[dict]] = []
    for item in items:
        if grouped and abs(item["y"] - statistics.mean(i["y"] for i in grouped[-1])) <= tolerance:
            grouped[-1].append(item)
        else:
            grouped.append([item])
    return [" ".join(i["text"] for i in sorted(group, key=lambda i: i["x"])) for group in grouped]


# =====================================================
# Invoice header (supplier, invoice number, date)
# =====================================================

GSTIN_RE = re.compile(r"\b\d{2}[A-Z]{5}\d{4}[A-Z][A-Z\d]Z[A-Z\d]\b")
INVOICE_NO_RE = re.compile(
    r"(?i)\b(?:tax\s+)?(?:invoice|inv|bill|challan)\s*(?:no|number|num|#)?\s*\.?\s*[:\-]?\s*([A-Za-z0-9][A-Za-z0-9/\-_.]{2,})"
)
DATE_RE = re.compile(r"(?i)\b(?:invoice\s*)?(?:dated?|date)\s*[:\-]?\s*(\d{1,2}[-/.]\d{1,2}[-/.]\d{2,4})")


def extract_invoice_meta(text_rows: list[str]) -> dict:
    meta = {"invoice_no": None, "invoice_date": None, "supplier_name": None, "supplier_gstin": None}
    for text in text_rows[:25]:
        text = _clean(text)
        if not meta["supplier_gstin"]:
            gstin = GSTIN_RE.search(text.upper())
            if gstin:
                meta["supplier_gstin"] = gstin.group(0)
        if not meta["invoice_no"]:
            match = INVOICE_NO_RE.search(text)
            if match and re.search(r"\d", match.group(1)):
                meta["invoice_no"] = match.group(1).strip(".,;:")
        if not meta["invoice_date"]:
            match = DATE_RE.search(text)
            if match:
                parsed = parse_expiry(match.group(1))
                if parsed:
                    meta["invoice_date"] = parsed.isoformat()
        if not meta["supplier_name"] and re.search(r"[A-Za-z]{4,}", text):
            if not re.search(
                r"(?i)\b(invoice|gstin|original|copy|page|duplicate|triplicate|tel|phone|email|d\.?l|licen[cs]e|address|patient|customer|doctor)\b",
                text,
            ) and not GSTIN_RE.search(text.upper()):
                meta["supplier_name"] = text[:80]
    return meta